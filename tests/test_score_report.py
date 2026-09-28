"""Checks for the PSH score-summary workbook exported from the results page."""
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from data_exporting import OutputSheet
from examinee_list import ExamineeTable
from grid_info import Field
from score_report import export_score_report
from scoring import score_results


class ScoreReportTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.destination = Path(directory.name) / "report.xlsx"

    def load_report(self):
        workbook = load_workbook(self.destination)
        self.addCleanup(workbook.close)
        return workbook.active

    def make_results(self):
        roster = ExamineeTable(["Examinee ID", "NAME"],
                               [["0001", "=Example Name"], ["0002", "No Scan Name"]]).bind()
        scans = OutputSheet([Field.STUDENT_ID, Field.IMAGE_FILE], 3)
        scans.add({Field.STUDENT_ID: "0001", Field.IMAGE_FILE: "one.png"}, ["A", "B", "C"])
        scans.add({Field.STUDENT_ID: "0003", Field.IMAGE_FILE: "three.png"}, ["D", "D", "C"])
        keys = OutputSheet([Field.IMAGE_FILE], 3)
        keys.add({}, ["A", "B", "D"])
        processed, _, _ = roster.merge(scans)
        review, _, _ = roster.merge(scans, include_unscanned=True)
        return review, score_results(processed, keys, 3)

    def test_summary_layout_names_scores_highlights_and_view_order(self):
        results, scores = self.make_results()
        export_score_report(results, scores, self.destination, row_order=[2, 0, 1])
        sheet = self.load_report()
        reference_path = Path(__file__).parent / "fixtures" / "reports" / "examinee-scores-layout.xlsx"
        reference = load_workbook(reference_path)
        self.addCleanup(reference.close)
        self.assertEqual(sheet.title, reference.active.title)
        self.assertEqual([sheet.cell(4, i).value for i in range(1, 4)],
                         [reference.active.cell(4, i).value for i in range(1, 4)])
        self.assertEqual(sheet["D4"].value, "Points (out of 3)")
        self.assertEqual(sheet.freeze_panes, "A5")
        self.assertEqual(sheet.auto_filter.ref, "A4:D7")
        self.assertEqual(sheet["A4"].fill.fgColor.rgb, "FF17324D")
        self.assertEqual(sheet.column_dimensions["B"].width, 48)
        self.assertEqual([sheet.cell(row, 1).value for row in (5, 6, 7)], ["0002", "0001", "0003"])
        self.assertEqual(sheet["A6"].data_type, "s")
        self.assertEqual(sheet["B6"].value, "=Example Name")
        self.assertEqual(sheet["B6"].data_type, "s")
        self.assertEqual(sheet["C6"].value, 66.67)
        self.assertEqual(sheet["D6"].value, 2)
        self.assertEqual(sheet["C6"].data_type, "n")
        self.assertEqual(sheet["C6"].number_format, "0.00")
        self.assertEqual(sheet["D6"].number_format, "0")
        self.assertEqual(sheet["B5"].value, "No Scan Name")
        self.assertIsNone(sheet["C5"].value)
        self.assertIsNone(sheet["D5"].value)
        for cell in sheet[5][:4]:
            self.assertEqual(cell.fill.fgColor.rgb, "FFFFF2D4")
        self.assertEqual(sheet["C7"].value, 0)
        self.assertEqual(sheet["D7"].value, 0)
        self.assertIsNone(sheet["C7"].fill.patternType)
        self.assertIn("no processed scan", sheet["F5"].value)
        self.assertIn("no matching examinee-list entry", sheet["F6"].value)
        self.assertEqual([sheet.cell(row, 2).value for row in (5, 6, 7)],
                         ["No Scan Name", "=Example Name", None])

    def test_no_answer_key_does_not_export_zero_scores(self):
        results, _ = self.make_results()
        export_score_report(results, None, self.destination)
        sheet = self.load_report()
        for row in range(5, 8):
            self.assertIsNone(sheet.cell(row, 3).value)
            self.assertIsNone(sheet.cell(row, 4).value)
            self.assertEqual(sheet.cell(row, 1).fill.fgColor.rgb, "FFFFF2D4")
        self.assertIn("2 processed scan(s) have no score", sheet["F6"].value)

    def test_repeated_examinee_id_keeps_each_scans_score_and_separate_name_parts(self):
        results = OutputSheet([Field.STUDENT_ID, Field.IMAGE_FILE, "List: First Name", "List: Last Name"], 2)
        results.add({Field.STUDENT_ID: "001", Field.IMAGE_FILE: "one.png",
                     "List: First Name": "Ana", "List: Last Name": "Example"}, ["A", "B"])
        results.add({Field.STUDENT_ID: "001", Field.IMAGE_FILE: "two.png",
                     "List: First Name": "Ana", "List: Last Name": "Example"}, ["C", "D"])
        keys = OutputSheet([Field.IMAGE_FILE], 2)
        keys.add({}, ["A", "B"])
        export_score_report(results, score_results(results, keys, 2), self.destination)
        sheet = self.load_report()
        self.assertEqual(sheet["B5"].value, "Ana Example")
        self.assertEqual(sheet["C5"].value, 100)
        self.assertEqual(sheet["C6"].value, 0)

    def test_empty_results_export_without_reference_records(self):
        results = OutputSheet([Field.STUDENT_ID, Field.IMAGE_FILE], 150)
        export_score_report(results, None, self.destination)
        sheet = self.load_report()
        self.assertEqual(sheet["D4"].value, "Points (out of 150)")
        self.assertEqual(sheet["F5"].value, "No examinees to report.")
        self.assertIsNone(sheet["A5"].value)
        self.assertIsNone(sheet["B5"].value)

    def test_export_requires_all_results_rows(self):
        results, scores = self.make_results()
        for order in ([0], [0, 0, 2], [0, 1, 3]):
            with self.subTest(order=order):
                with self.assertRaisesRegex(ValueError, "each results-page row"):
                    export_score_report(results, scores, self.destination, row_order=order)
        self.assertFalse(self.destination.exists())


if __name__ == "__main__":
    unittest.main()
