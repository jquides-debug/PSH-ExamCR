"""Regression checks for roster matching and the checked-results pipeline."""
import csv
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
import examinee_list
import grid_info
import process_input
import scoring
from data_exporting import OutputSheet
from grid_info import Field
from mcta_processing import create_answers_files


class ExamineeListTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name)

    def test_load_csv_with_bom_and_normalized_header(self):
        path = self.folder / "list.csv"
        path.write_text(' examinee_id ,Name,Department\n0001,"Peña, Ana",Medicine\n',
                        encoding="utf-8-sig")
        table = examinee_list.load_table(path)
        self.assertEqual(examinee_list.normalize_heading(table.headers[0]), "examineeid")
        roster = table.bind()
        self.assertEqual(roster.by_id["1"], ["0001", "Peña, Ana", "Medicine"])

    def test_xlsx_requires_id_header_and_uses_first_sheet(self):
        path = self.folder / "list.xlsx"
        workbook = Workbook()
        workbook.active.append(["#", "NAME", "EMAIL ADDRESS:"])
        workbook.active.append([1, "Test examinee", "test@example.test"])
        workbook.create_sheet("Other").append(["Ignored"])
        workbook.save(path)
        workbook.close()
        table = examinee_list.load_table(path)
        with self.assertRaisesRegex(ValueError, "No Examinee ID column"):
            table.bind()
        table.headers[0] = "Examinee ID"
        roster = table.bind()
        self.assertEqual(roster.by_id["1"][1], "Test examinee")

    def test_detects_header_variants_in_any_column(self):
        for heading in ("Examinee ID", "EXAMINEE ID", "examinee_id", "Examinee-ID",
                        "examineeid", "  Examinee  ID  "):
            with self.subTest(heading=heading):
                roster = examinee_list.ExamineeTable(
                    ["Name", heading, "Email"], [["Ana", "001", "ana@example.test"]]).bind()
                self.assertEqual(roster.id_index, 1)
                self.assertEqual(roster.by_id["1"][0], "Ana")

    def test_rejects_missing_and_ambiguous_id_headers(self):
        for headers, message in ((["#", "Name"], "No Examinee ID column"),
                                 (["Examinee ID", "examinee_id"], "exactly one")):
            with self.subTest(headers=headers):
                with self.assertRaisesRegex(ValueError, message):
                    examinee_list.ExamineeTable(headers, [["1", "Ana"]]).bind()

    def test_supplied_templates_can_be_completed_and_imported(self):
        templates = Path(__file__).resolve().parents[1] / "src" / "assets" / "templates"
        table = examinee_list.load_table(templates / "examinee-list.xlsx")
        self.assertEqual(table.headers, ["Examinee-ID", "NAME", "EMAIL ADDRESS:"])
        # Fill a copy in memory without modifying the user's template.
        table.rows = [["0000000001", "Example Examinee", "test@example.test"]]
        self.assertEqual(table.bind().by_id["1"][0], "0000000001")
        template_path = templates / "answer-key.csv"
        self.assertFalse(scoring.verify_answer_key_sheet(template_path))
        with template_path.open(encoding="utf-8-sig", newline="") as source:
            rows = list(csv.reader(source))
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0], ["Answer Key"] + [f"Q{i}" for i in range(1, 151)])
        key_path = self.folder / "completed_key.csv"
        with key_path.open("w", encoding="utf-8", newline="") as output:
            csv.writer(output).writerows([rows[0], [""] + ["A"] * 150])
        self.assertTrue(scoring.verify_answer_key_sheet(key_path))
        keys = OutputSheet([Field.IMAGE_FILE], 150)
        keys.add_file(key_path)
        self.assertEqual(keys.row_count, 1)
        self.assertEqual(keys.data[0][1:], [f"Q{i}" for i in range(1, 151)])
        self.assertEqual(len(keys.data[1][1:]), 150)
        self.assertTrue(set(keys.data[1][1:]) <= set("ABCDE"))
        self.assertTrue((templates / "examination-sheet.png").is_file())

    def test_invalid_rosters_do_not_silently_match(self):
        for body, message in (
            ("", "empty"),
            ("ID,,Name\n1,x,A", "heading"),
            ("Examinee ID,examinee_id\n1,2", "unique"),
            ("Examinee ID,Name\n", "no examinees"),
            ("Examinee ID,Name\n,A", "blank ID"),
            ("Examinee ID,Name\n001,A\n1,B", "duplicate ID"),
        ):
            with self.subTest(body=body):
                path = self.folder / "invalid.csv"
                path.write_text(body, encoding="utf-8")
                with self.assertRaisesRegex(ValueError, message):
                    examinee_list.load_table(path).bind()
        path = self.folder / "invalid.xlsx"
        path.write_text("not a workbook", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "Unable to read"):
            examinee_list.load_table(path)

    def test_merge_preserves_scans_and_protects_reserved_columns(self):
        roster = examinee_list.ExamineeTable(
            ["Examinee ID", "NAME", "Source File", "Q1", "Total Score (%)"],
            [["1", "Ana", "list-source", "F", "100"],
             ["2", "Ben", "", "", ""]]).bind()
        answers = OutputSheet([Field.STUDENT_ID, Field.IMAGE_FILE], 2)
        for identifier, source in (("0001", "a.png"), ("0001", "b.png"),
                                   ("0003", "c.png"), ("", "d.png")):
            answers.add({Field.STUDENT_ID: identifier, Field.IMAGE_FILE: source}, ["A", "[A|B]"])
        merged, matched, without_scan = roster.merge(answers)
        self.assertEqual((merged.row_count, matched, without_scan), (4, 2, 1))
        self.assertEqual(merged.data[1],
                         ["0001", "a.png", "Ana", "list-source", "F", "100", "Matched", "A", "[A|B]"])
        self.assertEqual(merged.data[3][2:7], ["", "", "", "", "Not in list"])
        self.assertEqual(merged.data[4][6], "Not in list")
        self.assertEqual(merged.data[0][3:6], ["List: Source File", "List: Q1", "List: Total Score (%)"])
        self.assertEqual(answers.data[0], ["Examinee ID", "Source File", "Q1", "Q2"])
        review, matched, without_scan = roster.merge(answers, include_unscanned=True)
        self.assertEqual((review.row_count, matched, without_scan), (5, 2, 1))
        self.assertEqual(review.data[:-1], merged.data)
        self.assertEqual(review.data[-1],
                         ["2", "", "Ben", "", "", "", "No scan detected", "", ""])
        self.assertEqual(merged.row_count, 4)

    def test_scoring_sorting_excel_and_mcta_with_imported_details(self):
        roster = examinee_list.ExamineeTable(
            ["Examinee ID", "Name", "Notes"],
            [["1", "Peña", "=literal"], ["2", "Another", "F"]]).bind()
        answers = OutputSheet([Field.STUDENT_ID, Field.IMAGE_FILE], 2)
        answers.add({Field.STUDENT_ID: "0002", Field.IMAGE_FILE: "b.png"}, ["B", "B"])
        answers.add({Field.STUDENT_ID: "0001", Field.IMAGE_FILE: "a.png"}, ["A", "[A|B]"])
        merged, _, _ = roster.merge(answers)
        merged.save(self.folder, "results", True, None)
        keys = OutputSheet([Field.IMAGE_FILE], 2)
        keys.add({}, ["A", "B"])
        scores = scoring.score_results(merged, keys, 2)
        scores.save(self.folder, "scores", True, None)
        self.assertEqual(scores.data[1][0:4], ["0001", "a.png", "Peña", "=literal"])
        self.assertEqual(scores.data[1][scores.data[0].index("Total Score (%)")], "50.0")
        self.assertEqual(scores.data[1][scores.first_question_column_index:], ["1", "0"])
        workbook = load_workbook(merged.save_excel(self.folder, None))
        self.addCleanup(workbook.close)
        sheet = workbook.active
        self.assertEqual(sheet["A2"].value, "0001")
        self.assertEqual(sheet["D2"].data_type, "s")
        self.assertEqual(sheet["D2"].value, "=literal")
        self.assertIsNone(sheet["D3"].fill.patternType)
        self.assertEqual(sheet["G2"].fill.fgColor.rgb, "FFFFC7CE")
        self.assertEqual(sheet.freeze_panes, "F2")
        with (self.folder / "results.csv").open(encoding="utf-8", newline="") as source:
            self.assertEqual(list(csv.reader(source))[1][2], "Peña")
        create_answers_files(merged, self.folder, None)
        with (self.folder / "mcta_exam_results.csv").open(encoding="utf-8", newline="") as source:
            self.assertEqual(list(csv.reader(source)),
                             [["", "Q1", "Q2"], ["Student0", "A", "[A|B]"], ["Student1", "B", "B"]])

    def test_empty_processed_batch_reports_all_list_entries_without_scan(self):
        roster = examinee_list.ExamineeTable(["Examinee ID", "Name"], [["1", "Ana"]]).bind()
        answers = OutputSheet([Field.STUDENT_ID, Field.IMAGE_FILE], 2)
        answers.clean_up()
        merged, matched, without_scan = roster.merge(answers)
        self.assertEqual((merged.row_count, matched, without_scan), (0, 0, 1))
        self.assertEqual(merged.data[0][-2:], ["Q1", "Q2"])
        review, matched, without_scan = roster.merge(answers, include_unscanned=True)
        self.assertEqual((review.row_count, matched, without_scan), (1, 0, 1))
        self.assertEqual(review.data[1], ["1", "", "Ana", "No scan detected", "", ""])

    def test_processing_pipeline_exports_and_displays_matched_details(self):
        roster = examinee_list.ExamineeTable(
            ["Examinee ID", "Name"], [["1", "Ana"], ["2", "Ben"]]).bind()
        key_file = self.folder / "key.csv"
        key_file.write_text("Q1,Q2\nA,B\n", encoding="utf-8")
        form = SimpleNamespace(fields={Field.STUDENT_ID: grid_info.form_150q.fields[Field.STUDENT_ID]},
                               questions=grid_info.form_150q.questions[:2], num_questions=2)
        from unittest.mock import Mock
        tracker = Mock()
        with patch.object(process_input.image_utils, "get_image"), \
                patch.object(process_input.image_utils, "prepare_scan_for_processing"), \
                patch.object(process_input.image_utils, "dilate"), \
                patch.object(process_input.corner_finding, "find_corner_marks"), \
                patch.object(process_input.grid_r, "Grid"), \
                patch.object(process_input.grid_r, "get_group_from_info"), \
                patch.object(process_input.grid_r, "calculate_bubble_fill_threshold"), \
                patch.object(process_input.grid_r, "read_field_as_string", return_value="0000000001"), \
                patch.object(process_input.grid_r, "read_answer_as_string", side_effect=["A", "[A|B]"]):
            process_input.process_input([Path("scan.png")], self.folder, False, False,
                                        key_file, None, True, True, True, form, tracker,
                                        None, examinees=roster)
        tracker.show_results.assert_called_once()
        results, scores, workbook_path, rejected, _, summary = tracker.show_results.call_args.args
        self.assertEqual(results.data[1][2:4], ["Ana", "Matched"])
        self.assertEqual(scores.data[1][2:4], ["Ana", "Matched"])
        self.assertEqual(results.row_count, 2)
        self.assertEqual(results.data[2], ["2", "", "Ben", "No scan detected", "", ""])
        self.assertEqual(scores.row_count, 1)
        self.assertEqual(rejected, 0)
        self.assertIn("1 matched scans", summary)
        self.assertIn("1 list entries without a processed scan", summary)
        self.assertTrue(workbook_path.exists())
        for name in ("results.csv", "scores.csv", "keys.csv", "mcta_exam_results.csv"):
            self.assertTrue((self.folder / name).exists(), name)
            with (self.folder / name).open(encoding="utf-8", newline="") as source:
                self.assertEqual(len(list(csv.reader(source))), 2)


if __name__ == "__main__":
    unittest.main()
