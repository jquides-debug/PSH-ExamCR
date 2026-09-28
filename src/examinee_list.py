"""Load an examinee roster and attach its details to processed answer sheets."""
import csv
from dataclasses import dataclass
from pathlib import Path
import re
import typing as tp
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from data_exporting import COLUMN_NAMES, OutputSheet
from grid_info import Field


MATCH_COLUMN = "Examinee list match"
NO_SCAN_STATUS = "No scan detected"


def normalize_heading(value: str) -> str:
    return re.sub(r"[\s_-]+", "", value).casefold()


def normalize_id(value: str) -> str:
    """Match numeric Excel IDs to zero-padded bubble IDs without changing exports."""
    value = value.strip()
    if re.fullmatch(r"[0-9]+", value):
        return value.lstrip("0") or "0"
    return value


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


@dataclass
class ExamineeTable:
    headers: tp.List[str]
    rows: tp.List[tp.List[str]]

    def bind(self):
        candidates = [index for index, header in enumerate(self.headers)
                      if normalize_heading(header) == "examineeid"]
        if not candidates:
            raise ValueError(
                "No Examinee ID column found. Use the examinee-list template or rename "
                "the column containing the sheet IDs to 'Examinee ID'.")
        if len(candidates) != 1:
            raise ValueError("Keep exactly one Examinee ID column in the examinee list.")
        id_index = candidates[0]
        by_id = {}
        for row_number, row in enumerate(self.rows, 2):
            if not any(row):
                continue
            key = normalize_id(row[id_index])
            if not key:
                raise ValueError(f"Examinee list row {row_number} has a blank ID.")
            if key in by_id:
                raise ValueError(
                    f"Examinee list has duplicate ID '{row[id_index]}' at row {row_number}. "
                    "Each ID must occur once (leading zeroes are ignored when matching).")
            by_id[key] = row
        if not by_id:
            raise ValueError("The examinee list has no examinees.")
        return ExamineeList(self.headers, id_index, by_id)


def load_table(path: Path) -> ExamineeTable:
    """Read headers on row 1; XLSX imports the first worksheet."""
    try:
        if path.suffix.lower() == ".csv":
            with path.open(encoding="utf-8-sig", newline="") as source:
                rows = [[cell_text(value) for value in row] for row in csv.reader(source)]
        elif path.suffix.lower() == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                rows = [[cell_text(value) for value in row]
                        for row in workbook.worksheets[0].iter_rows(values_only=True)]
            finally:
                workbook.close()
        else:
            raise ValueError("Choose an Excel (.xlsx) or UTF-8 CSV (.csv) examinee list.")
    except (OSError, UnicodeError, csv.Error, BadZipFile, InvalidFileException,
            KeyError, IndexError) as error:
        raise ValueError(f"Unable to read the examinee list: {error}") from error

    width = max((index + 1 for row in rows for index, value in enumerate(row) if value),
                default=0)
    if not width:
        raise ValueError("The examinee list is empty.")
    rows = [(row + [""] * max(0, width - len(row)))[:width] for row in rows]
    headers = rows[0]
    if not all(headers):
        raise ValueError("Every examinee-list column must have a heading in the first row.")
    if len({normalize_heading(header) for header in headers}) != len(headers):
        raise ValueError("Examinee-list column headings must be unique.")
    return ExamineeTable(headers, rows[1:])


@dataclass
class ExamineeList:
    headers: tp.List[str]
    id_index: int
    by_id: tp.Dict[str, tp.List[str]]

    def merge(self, results: OutputSheet, include_unscanned: bool = False):
        """Enrich scans, optionally adding roster-only rows for the results page."""
        detail_indices = [index for index in range(len(self.headers)) if index != self.id_index]
        reserved = {normalize_heading(name) for name in
                    list(COLUMN_NAMES.values()) + results.data[0] + [MATCH_COLUMN]}
        names = []
        for index in detail_indices:
            name = self.headers[index]
            while normalize_heading(name) in reserved or re.fullmatch(r"Q\d+", name, re.I):
                name = "List: " + name
            names.append(name)
            reserved.add(normalize_heading(name))

        merged = OutputSheet(results.field_columns + names + [MATCH_COLUMN],
                             results.num_questions)
        start = results.first_question_column_index
        merged.data = [results.data[0][:start] + names + [MATCH_COLUMN]
                       + results.data[0][start:]]
        id_index = results.field_columns.index(Field.STUDENT_ID)
        matched_ids = set()
        matched_count = 0
        for row in results.data[1:]:
            key = normalize_id(row[id_index])
            examinee = self.by_id.get(key)
            if examinee is not None:
                details = [examinee[index] for index in detail_indices]
                matched_count += 1
                matched_ids.add(key)
            else:
                details = [""] * len(detail_indices)
            merged.data.append(row[:start] + details
                               + ["Matched" if examinee is not None else "Not in list"]
                               + row[start:])
        if include_unscanned:
            for key, examinee in self.by_id.items():
                if key in matched_ids:
                    continue
                fields = [""] * start
                fields[id_index] = examinee[self.id_index]
                details = [examinee[index] for index in detail_indices]
                merged.data.append(fields + details + [NO_SCAN_STATUS]
                                   + [""] * (len(results.data[0]) - start))
        merged.row_count = len(merged.data) - 1
        return merged, matched_count, len(self.by_id) - len(matched_ids)
