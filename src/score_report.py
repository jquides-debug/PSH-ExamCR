"""Export the results-page roster in the PSH examinee-score report layout."""
from datetime import datetime
import math
from pathlib import Path
import typing as tp

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from data_exporting import OutputSheet
from examinee_list import MATCH_COLUMN, NO_SCAN_STATUS, normalize_heading


def _number(value: str) -> tp.Optional[float]:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def export_score_report(results: OutputSheet, scores: tp.Optional[OutputSheet],
                        destination: Path, row_order: tp.Optional[tp.Sequence[int]] = None):
    """Write a fresh workbook, keeping reference-workbook records out of exports."""
    headings = results.data[0]
    id_index = headings.index("Examinee ID")
    source_index = headings.index("Source File")
    match_index = headings.index(MATCH_COLUMN) if MATCH_COLUMN in headings else None
    name_columns = {}
    for index, heading in enumerate(headings[:results.first_question_column_index]):
        # Imported name columns can be prefixed to protect the scan's field names.
        while heading.startswith("List: "):
            heading = heading[len("List: "):]
        name_columns[normalize_heading(heading.rstrip(":"))] = index

    def examinee_name(row):
        for heading in ("name", "examineename", "fullname"):
            index = name_columns.get(heading)
            if index is not None and row[index].strip():
                return row[index]
        return " ".join(row[name_columns[part]].strip()
                        for part in ("firstname", "middlename", "lastname")
                        if part in name_columns and row[name_columns[part]].strip())

    score_lookup = {}
    question_count = results.num_questions
    if scores is not None:
        score_headings = scores.data[0]
        score_id = score_headings.index("Examinee ID")
        score_source = score_headings.index("Source File")
        percent_index = score_headings.index("Total Score (%)")
        points_index = score_headings.index("Total Points")
        for row in scores.data[1:]:
            score_lookup[(row[score_id], row[score_source])] = (
                _number(row[percent_index]), _number(row[points_index]))
        if scores.row_count:
            question_count = max(len(row) - scores.first_question_column_index
                                 for row in scores.data[1:])

    order = list(range(results.row_count)) if row_order is None else list(row_order)
    if sorted(order) != list(range(results.row_count)):
        raise ValueError("The export must include each results-page row exactly once.")

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Examinee scores"
    sheet["A2"] = f"{datetime.now().year} PSH Certifying Examination - Examinee Scores"
    sheet["A2"].font = Font(name="Arial", size=14, bold=True)
    header_fill = PatternFill(fill_type="solid", fgColor="FF17324D")
    missing_fill = PatternFill(fill_type="solid", fgColor="FFFFF2D4")
    for column, title in enumerate(("Examinee No.", "Examinee Name", "Score (%)",
                                    f"Points (out of {question_count})"), 1):
        cell = sheet.cell(4, column, title)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
        cell.fill = header_fill
    sheet["F4"] = "Matching note"
    sheet["F4"].font = Font(name="Arial", size=10, bold=True)
    sheet.row_dimensions[4].height = 28

    missing_scans = unscored = unmatched = 0
    for excel_row, result_index in enumerate(order, 5):
        row = results.data[result_index + 1]
        status = row[match_index] if match_index is not None else ""
        no_scan = status == NO_SCAN_STATUS
        percent, points = (None, None) if no_scan else score_lookup.get(
            (row[id_index], row[source_index]), (None, None))
        missing_scans += int(no_scan)
        unscored += int(not no_scan and percent is None)
        unmatched += int(status == "Not in list")
        for column, value in enumerate((row[id_index], examinee_name(row), percent, points), 1):
            cell = sheet.cell(excel_row, column, value)
            cell.font = Font(name="Arial", size=10)
            if column <= 2:
                # Keep leading zeroes and names starting with '=' as literal text.
                cell.data_type = "s"
            cell.number_format = "@" if column == 1 else "0.00" if column == 3 else "0" if column == 4 else "General"
            if percent is None:
                cell.fill = missing_fill
        sheet.row_dimensions[excel_row].height = 22

    notes = []
    if missing_scans:
        notes.append(f"{missing_scans} examinee(s) have no processed scan; scores are blank.")
    if unscored:
        notes.append(f"{unscored} processed scan(s) have no score; no zero score was assigned.")
    if unmatched:
        notes.append(f"{unmatched} processed scan(s) have no matching examinee-list entry.")
    if not notes:
        notes.append("All listed examinees have scores." if results.row_count else "No examinees to report.")
    notes.append("Highlighted rows have no score available.")
    for note_row, note in enumerate(notes, 5):
        sheet.cell(note_row, 6, note).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.row_dimensions[note_row].height = max(sheet.row_dimensions[note_row].height or 22, 32)

    for column, width in (("A", 17), ("B", 48), ("C", 17), ("D", 23), ("E", 3), ("F", 58)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:D{max(4, results.row_count + 4)}"
    try:
        workbook.save(destination)
    finally:
        workbook.close()
    return destination
